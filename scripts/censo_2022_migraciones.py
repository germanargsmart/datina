"""
Censo 2022: lugar de nacimiento de la población por departamento (etapa 26, octubre de 2026).
El INDEC publica los resultados de migraciones de cada provincia en su página (censo.gob.ar/index.php/datos_definitivos_<provincia>/):
el cuadro 1 trae, por departamento, la población en viviendas particulares nacida en la provincia, en otra provincia y en otro país, y el
cuadro 3, una hoja por departamento con la población nacida en otro país según el país o la región de nacimiento. Como el censo no cambia,
este script se corre una sola vez: lee los enlaces de las 24 páginas, baja los 48 cuadros a memoria y guarda solo el resultado en
datos/manual/censo2022_migraciones_departamentos.csv (una fila por departamento, con cantidades: los porcentajes los calcula 02_procesar.py).
Las hojas del cuadro 3 no traen el código del departamento: se unen al cuadro 1 por el nombre, sin tildes ni signos.
Uso: python scripts/censo_2022_migraciones.py
"""
import io, json, pathlib, re, unicodedata, urllib.request
import openpyxl
import pandas as pd

RAIZ = pathlib.Path(__file__).resolve().parents[1]
SALIDA = RAIZ / "datos" / "manual" / "censo2022_migraciones_departamentos.csv"
UA = {"User-Agent": "Mozilla/5.0 (argentina-en-datos)"}
PAGINAS = ["bsas", "caba", "catamarca", "chaco", "chubut", "cordoba", "corrientes", "entre_rios", "formosa", "jujuy", "lapampa", "larioja",
           "mendoza", "misiones", "neuquen", "rionegro", "salta", "sanjuan", "sanluis", "santa_cruz", "santafe", "santiago_del_estero", "tdf",
           "tucuman"]
# filas del cuadro 3 que se guardan (los países y las regiones, sin los subtotales de continente) -> columna de salida
PAISES = {"bolivia": "bolivia", "brasil": "brasil", "chile": "chile", "paraguay": "paraguay", "uruguay": "uruguay", "colombia": "colombia",
          "peru": "peru", "venezuela": "venezuela", "resto de america": "resto_america", "espana": "espana", "italia": "italia",
          "resto de europa": "resto_europa", "asia": "asia", "africa": "africa", "oceania": "oceania", "ignorado": "ignorado"}
ALIAS = {"cerrilos": "cerrillos"}  # Salta: el cuadro 3 escribe "Cerrilos"


def leer(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=300) as r:
        return r.read()


def limpio(v):
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip() if v is not None else ""


def norm(s):
    s = limpio(s).replace("°", "o").replace("º", "o")
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower())


def numero(v):
    """Los cuadros escriben "-" para el cero absoluto."""
    return 0 if limpio(v) in ("-", "") else int(v)


def codigo(v):
    s = limpio(v)
    if not re.fullmatch(r"\d+(\.0)?", s):
        return None
    s = s.split(".")[0]
    return s.zfill(2) if len(s) <= 2 else s.zfill(5)  # la Ciudad escribe 2007 y no 02007


def libro(url):
    b = leer(url)
    assert b[:2] == b"PK", f"{url} no es una planilla"
    return openpyxl.load_workbook(io.BytesIO(b), data_only=True, read_only=False)


filas = []
for pagina in PAGINAS:
    html = leer(f"https://censo.gob.ar/index.php/datos_definitivos_{pagina}/").decode("utf-8", "ignore")
    enlaces = sorted(set(re.findall(r'href="([^"]+_migraciones_c[13]_\d+\.xlsx)"', html)))
    c1 = [u for u in enlaces if "_migraciones_c1_" in u]
    c3 = [u for u in enlaces if "_migraciones_c3_" in u]
    assert len(c1) == 1 and len(c3) == 1, (pagina, enlaces)
    # cuadro 1: la primera hoja que no es el índice (Buenos Aires tiene además una hoja "bis" con el Gran Buenos Aires, que no se usa)
    wb = libro(c1[0])
    ws = [w for w in wb.worksheets if not limpio(w.title).lower().startswith("índice") and not limpio(w.title).endswith("bis")][0]
    deps = []
    for r in range(5, ws.max_row + 1):
        c = codigo(ws.cell(r, 1).value)
        if c is None:
            continue  # títulos, notas y los subtotales de Buenos Aires, que no tienen código
        vals = [numero(ws.cell(r, k).value) for k in range(3, 7)]
        deps.append({"codigo_dep": c, "departamento": limpio(ws.cell(r, 2).value), "poblacion_vp": vals[0], "en_la_provincia": vals[1],
                     "en_otra_provincia": vals[2], "en_otro_pais": vals[3]})
    prov, deps = deps[0], deps[1:]
    assert len(prov["codigo_dep"]) == 2 and all(d["codigo_dep"][:2] == prov["codigo_dep"] for d in deps), pagina
    for k in ("poblacion_vp", "en_la_provincia", "en_otra_provincia", "en_otro_pais"):
        assert sum(d[k] for d in deps) == prov[k], (pagina, k)
    assert all(d["poblacion_vp"] == d["en_la_provincia"] + d["en_otra_provincia"] + d["en_otro_pais"] for d in deps), pagina
    # cuadro 3: una hoja por departamento, con el nombre en el título
    por_nombre = {norm(d["departamento"]): d for d in deps}
    wb = libro(c3[0])
    hojas = [w for w in wb.worksheets if limpio(w.title).lower().startswith("cuadro")][1:]
    assert len(hojas) == len(deps), (pagina, len(hojas), len(deps))
    for ws in hojas:
        m = re.search(r"(departamento|partido|comuna)\s+(.*?)\.\s+Poblaci", limpio(ws.cell(2, 1).value), re.I)
        nombre = ("Comuna " if m.group(1).lower() == "comuna" else "") + m.group(2)
        d = por_nombre[ALIAS.get(norm(nombre), norm(nombre))]
        assert "total_c3" not in d, (pagina, nombre)
        for r in range(6, ws.max_row + 1):
            etq = norm(re.sub(r"\(\d\)", "", limpio(ws.cell(r, 1).value)))
            etq = {norm(k): k for k in PAISES}.get(etq)
            if etq:
                d[PAISES[etq]] = numero(ws.cell(r, 2).value)
            elif limpio(ws.cell(r, 1).value) == "Total":
                d["total_c3"] = numero(ws.cell(r, 2).value)
        assert d["total_c3"] == d["en_otro_pais"] == sum(d[c] for c in PAISES.values()), (pagina, nombre)
    filas += deps
    print("listo", pagina, len(deps), "departamentos")

salida = pd.DataFrame(filas).drop(columns="total_c3")
salida.insert(1, "codigo_indec", salida.codigo_dep.str[:2])
# controles: los 527 departamentos del mapa del sitio y los totales del país que ya usa la sección Migraciones
geo = json.loads((RAIZ / "sitio" / "datos" / "departamentos.topo.json").read_text(encoding="utf-8"))
ids = {g["properties"]["id"] for g in list(geo["objects"].values())[0]["geometries"]}
assert len(salida) == 527 and set(salida.codigo_dep) == ids and salida.codigo_dep.is_unique
assert salida.poblacion_vp.sum() == 45618787 and salida.en_otro_pais.sum() == 1933463, (salida.poblacion_vp.sum(), salida.en_otro_pais.sum())
SALIDA.parent.mkdir(parents=True, exist_ok=True)
salida.to_csv(SALIDA, index=False)
print(f"{len(salida)} departamentos, {salida.en_otro_pais.sum():,} personas nacidas en otro país, {salida.ignorado.sum():,} sin país declarado")
